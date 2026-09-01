"""
aletheia/core/reporting/exporter.py

Generates PDF and Excel exports of agent run results.

PDF:  Jinja2 → HTML → WeasyPrint (if available) or raw HTML fallback
Excel: openpyxl workbook with multiple sheets
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from jinja2 import Environment, FileSystemLoader, select_autoescape

logger = logging.getLogger(__name__)

_TEMPLATE_DIR = Path(__file__).parent.parent / "templates"


def _build_jinja_env() -> Environment:
    return Environment(
        loader=FileSystemLoader(str(_TEMPLATE_DIR)),
        autoescape=select_autoescape(["html"]),
    )


def _flatten_run_data(run: dict[str, Any]) -> dict[str, Any]:
    """Extract and normalise fields from a raw run result dict."""
    result = run.get("result_json") or run.get("result") or {}
    if isinstance(result, str):
        import json
        try:
            result = json.loads(result)
        except Exception:
            result = {}

    # Holdings from portfolio snapshot in run request
    request = run.get("request") or {}
    portfolio = request.get("portfolio") or {}
    holdings = portfolio.get("holdings") or []
    portfolio_name = portfolio.get("name", "Unknown Portfolio")

    # Calculate portfolio value
    portfolio_value = sum(
        (h.get("quantity", 0) * h.get("average_price", 0)) for h in holdings
    )

    # Extract recommendations
    recommendations = result.get("recommendations") or []

    # Executive summary
    executive_summary = result.get("executive_summary") or result.get("summary", "")

    # Risk
    risk_raw = result.get("risk_metrics") or {}
    risk = None
    if risk_raw:
        class _Risk:
            def __init__(self, d: dict):
                self.var_95 = float(d.get("var_95", 0))
                self.sharpe = float(d.get("sharpe", 0))
                self.max_drawdown = float(d.get("max_drawdown", 0))
                self.win_rate = float(d.get("win_rate", 0.5))
        risk = _Risk(risk_raw)

    # Tax scenarios
    tax_scenarios_raw = result.get("tax_scenarios") or []

    # Overall verdict
    overall_verdict = result.get("final_verdict") or result.get("decision", "—")

    # Confidence
    confidence_raw = result.get("confidence", 0)
    if isinstance(confidence_raw, (int, float)):
        confidence = float(confidence_raw)
    else:
        confidence = 0.5

    return {
        "run_id": run.get("run_id", "unknown"),
        "status": run.get("status", "unknown").upper(),
        "generated_at": datetime.now(tz=timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "portfolio_name": portfolio_name,
        "portfolio_value": f"{portfolio_value:,.0f}",
        "num_holdings": len(holdings),
        "holdings": holdings,
        "recommendations": recommendations,
        "executive_summary": executive_summary,
        "risk": risk,
        "tax_scenarios": tax_scenarios_raw,
        "confidence": confidence,
        "overall_verdict": overall_verdict,
    }


def render_html(run: dict[str, Any]) -> str:
    """Render the report as an HTML string."""
    env = _build_jinja_env()
    template = env.get_template("report.html")
    ctx = _flatten_run_data(run)
    return template.render(**ctx)


def export_pdf(run: dict[str, Any]) -> bytes:
    """
    Export run result as a PDF binary.
    Uses WeasyPrint if available; falls back to returning the raw HTML bytes.
    """
    html_str = render_html(run)
    try:
        from weasyprint import HTML  # type: ignore[import-untyped]
        pdf_bytes: bytes = HTML(string=html_str).write_pdf()
        return pdf_bytes
    except ImportError:
        logger.warning(
            "weasyprint not installed — falling back to raw HTML export. "
            "Install with: pip install weasyprint"
        )
        return html_str.encode("utf-8")
    except Exception as exc:
        logger.error("WeasyPrint PDF rendering failed: %s — returning HTML", exc)
        return html_str.encode("utf-8")


def export_excel(run: dict[str, Any]) -> bytes:
    """Export run result as an Excel (.xlsx) binary."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel export: pip install openpyxl") from exc

    ctx = _flatten_run_data(run)
    wb = openpyxl.Workbook()

    # --- Sheet 1: Summary ---
    ws_summary = wb.active
    ws_summary.title = "Summary"  # type: ignore[attr-defined]

    PURPLE = "7C3AED"
    DARK = "1E1B4B"

    def header_row(ws, values: list[str], row: int) -> None:
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor=DARK)
            cell.alignment = Alignment(horizontal="center")

    def title_row(ws, text: str, row: int, ncols: int = 5) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor=PURPLE)
        cell.alignment = Alignment(horizontal="left")

    title_row(ws_summary, "ALETHEIA — Investment Intelligence Report", 1)

    summary_data = [
        ("Run ID", ctx["run_id"]),
        ("Status", ctx["status"]),
        ("Generated At", ctx["generated_at"]),
        ("Portfolio", ctx["portfolio_name"]),
        ("Portfolio Value (₹)", ctx["portfolio_value"]),
        ("Holdings", str(ctx["num_holdings"])),
        ("Confidence", f"{ctx['confidence']*100:.0f}%"),
        ("Verdict", ctx["overall_verdict"]),
    ]
    ws_summary.append([])  # blank row
    for label, value in summary_data:
        row_idx = ws_summary.max_row + 1
        ws_summary.append([label, value])
        ws_summary.cell(row=row_idx, column=1).font = Font(bold=True, color=PURPLE)

    # Executive summary
    ws_summary.append([])
    ws_summary.append(["Executive Summary"])
    ws_summary.cell(row=ws_summary.max_row, column=1).font = Font(bold=True, size=11)
    ws_summary.append([ctx.get("executive_summary", "N/A")])
    ws_summary.column_dimensions["A"].width = 25
    ws_summary.column_dimensions["B"].width = 80

    # --- Sheet 2: Holdings ---
    ws_h = wb.create_sheet("Holdings")
    title_row(ws_h, "Portfolio Holdings", 1, 6)
    ws_h.append([])
    header_row(ws_h, ["Symbol", "Exchange", "Quantity", "Avg Price (₹)", "Value (₹)", "Asset Type"], 3)
    for h in ctx["holdings"]:
        ws_h.append([
            h.get("symbol"), h.get("exchange"), h.get("quantity"),
            h.get("average_price"), h.get("quantity", 0) * h.get("average_price", 0),
            h.get("asset_type"),
        ])
    for col in range(1, 7):
        ws_h.column_dimensions[get_column_letter(col)].width = 18

    # --- Sheet 3: Recommendations ---
    ws_r = wb.create_sheet("Recommendations")
    title_row(ws_r, "Agent Recommendations", 1, 5)
    ws_r.append([])
    header_row(ws_r, ["Symbol", "Action", "Target Price", "Confidence", "Reasoning"], 3)
    for rec in ctx["recommendations"]:
        ws_r.append([
            rec.get("symbol"), rec.get("action"),
            rec.get("target_price"), f"{rec.get('confidence', 0)*100:.0f}%",
            (rec.get("reasoning") or "")[:500],
        ])
    ws_r.column_dimensions["E"].width = 60

    # --- Sheet 4: Tax Scenarios ---
    ws_t = wb.create_sheet("Tax Scenarios")
    title_row(ws_t, "Sage Tax Projections", 1, 5)
    ws_t.append([])
    header_row(ws_t, ["Scenario", "Duration", "Tax Rate (%)", "Est. Tax (₹)", "Net Return (₹)"], 3)
    for s in ctx.get("tax_scenarios") or []:
        ws_t.append([
            s.get("label"), s.get("duration"), s.get("tax_rate"),
            s.get("estimated_tax"), s.get("net_return"),
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
